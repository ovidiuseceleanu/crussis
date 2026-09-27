#!/usr/bin/env python3
"""Fișier de import BikeVerse pentru bicicletele Crussis Avinox."""

import csv
import html
import json
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path("/Users/ovidiuseceleanu/Documents/crussis")
IMG_ROOT = ROOT / "Avinox images"
SIZE_RE = re.compile(r"(XXL|XL|XS|S|M|L|\d{2})")
SIZE_ORDER = {"XXS": 0, "XS": 1, "S": 2, "M": 3, "L": 4, "XL": 5, "XXL": 6}

DISCIPLINE = {
    "Full suspension*Mountain": "ENDURO",
    "Mountain": "TRAIL",
    "SUV*City*Trek": "TREKKING",
    "City": "CITY",
    "SUV*Full suspension": "ENDURO",
    "City*Full suspension": "TREKKING",
    "Mountain*City*Full suspension": "ENDURO",
}

COLUMNS = [
    "category", "discipline", "brand", "family_name", "model_name", "model_color",
    "frame_size", "sku", "list_price_eur", "purchase_price_eur", "stock",
    "stock_status", "description_html", "description_html_en", "display", "ean", "notes", "image_files",
    "frame", "fork", "shock", "headset", "stem", "handlebar", "grips",
    "crankset", "rear_derailleur", "shifters_levers", "brakes", "cassette", "chain",
    "rim_front", "rim_rear", "hub_front", "hub_rear", "wheels",
    "tyre_front", "tyre_rear", "seatpost", "saddle", "pedals",
    "motor", "motor_brand", "motor_model", "battery", "battery_charger",
    "motor_power_max_w", "wheel_size_front", "wheel_size_rear",
]


def text(node, tag):
    if node is None:
        return ""
    value = node.findtext(tag)
    return (value or "").strip()


def unescape(value):
    value = html.unescape(value or "")
    value = value.replace("\xa0", " ")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def params_of(item):
    bag = defaultdict(list)
    for param in item.findall("PARAM"):
        name = (param.findtext("PARAM_NAME") or "").strip()
        value = unescape(param.findtext("VAL") or "")
        if name and value:
            bag[name].append(value)
    return bag


def first(bag, *names):
    for name in names:
        values = bag.get(name) or []
        if values:
            return values[0]
    return ""


def join_unique(parts):
    out = []
    for part in parts:
        part = (part or "").strip()
        if part and part not in out:
            out.append(part)
    return " | ".join(out)


def split_size(productno):
    match = re.search(rf" \({SIZE_RE.pattern}\)$", productno)
    if match:
        return productno[: match.start()].strip(), match.group(1)
    match = re.search(rf"\)({SIZE_RE.pattern})$", productno)
    if match:
        return productno[: match.start() + 1], match.group(1)
    return productno, ""


def as_model(name):
    name = name.strip()
    if name.startswith("ONE-"):
        return "e-" + name[4:]
    return name


def parse_eu_price(raw):
    raw = (raw or "").strip()
    if not raw:
        return None
    raw = raw.replace(".", "").replace(",", ".")
    return round(float(raw), 2)


def purchase_from_list(list_price):
    return round(float(list_price) / 1.21 / 1.35, 2)


def parse_pieces(raw):
    raw = re.sub(r"<[^>]+>", " ", raw or "")
    raw = re.sub(r"\s+", " ", raw).strip()
    plus = re.search(r"(\d+)\s*\+", raw)
    if plus:
        return int(plus.group(1))
    number = re.search(r"(\d+)", raw)
    return int(number.group(1)) if number else None


def clean_qty_label(raw):
    raw = re.sub(r"<[^>]+>", " ", raw or "")
    raw = re.sub(r"\s+", " ", raw).strip()
    raw = re.sub(r"^Order with\s*", "", raw, flags=re.I).strip()
    raw = re.sub(r"\s*pc$", "", raw, flags=re.I).strip()
    return raw


def production_windows(b2b):
    windows = []
    if not b2b:
        return windows
    if b2b.get("opts"):
        for opt in b2b["opts"]:
            if (opt.get("pid") or "") == "real_stock":
                continue
            label = clean_qty_label(opt.get("label") or opt.get("pid") or "")
            qty = clean_qty_label(opt.get("qty") or "")
            if not label:
                continue
            windows.append(f"{label}: {qty}" if qty else label)
    else:
        infos = b2b.get("info") or []
        qtys = b2b.get("ks") or []
        for index, info in enumerate(infos):
            label = clean_qty_label(info)
            if not label or label.lower() in {"in stock", "sold out", "select availability"}:
                continue
            qty = clean_qty_label(qtys[index]) if index < len(qtys) else ""
            windows.append(f"{label}: {qty}" if qty else label)
    deduped = []
    for window in windows:
        if window not in deduped:
            deduped.append(window)
    return deduped


def warehouse_qty(b2b):
    if not b2b:
        return None
    for opt in b2b.get("opts") or []:
        if (opt.get("pid") or "") == "real_stock":
            return parse_pieces(opt.get("qty") or "")
    infos = b2b.get("info") or []
    qtys = b2b.get("ks") or []
    for index, info in enumerate(infos):
        label = clean_qty_label(info).lower()
        if label == "in stock":
            qty = parse_pieces(qtys[index]) if index < len(qtys) else None
            return qty
    return None


def watts(performance):
    match = re.search(r"(\d+)\s*W", performance or "", flags=re.I)
    return int(match.group(1)) if match else ""


def range_text(raw):
    raw = (raw or "").strip()
    if not raw:
        return ""
    text_value = raw
    if re.search(r"\bup to\b", text_value, flags=re.I):
        text_value = re.sub(r"\bup to\b", "până la", text_value, flags=re.I)
    if "km" not in text_value.lower():
        text_value = text_value + " km"
    return "Autonomie: " + text_value


def load_text(raw):
    raw = re.sub(r"\s+", " ", raw or "").strip()
    raw = re.sub(r"\bkg\s+kg\b", "kg", raw, flags=re.I)
    return f"Sarcină maximă: {raw}" if raw else ""


def battery_text(bag):
    base = first(bag, "batteries")
    kind = first(bag, "Battery type")
    label = {"Removable": "detașabilă", "Fixed": "fixă"}.get(kind, "")
    if label and label not in base.lower() and kind.lower() not in base.lower():
        return f"{base}, baterie {label}" if base else f"baterie {label}"
    return base


def paragraphs_html(text):
    parts = []
    for para in (text or "").replace("\r", "").split("\n"):
        para = para.strip()
        if para:
            parts.append("<p>" + html.escape(para) + "</p>")
    return "\n".join(parts)


def image_urls(item):
    urls = []
    main = text(item, "IMGURL")
    if main:
        urls.append(main.split("?")[0])
    for node in item.findall("IMGURL_ALTERNATIVE"):
        url = (node.text or "").strip()
        if url:
            url = url.split("?")[0]
            if url not in urls:
                urls.append(url)
    return urls


def safe_part(value):
    value = value.replace("/", "-").replace("\\", "-")
    value = re.sub(r'[<>:"|?*]', "", value)
    value = re.sub(r"\s+", " ", value).strip().rstrip(".")
    return value or "necunoscut"


def extension_of(url):
    suffix = Path(url.split("?")[0]).suffix.lower()
    if suffix in {".jpg", ".jpeg", ".png", ".webp", ".gif"}:
        return suffix
    return ".jpg"


def load_stock():
    stock = {}
    root = ET.parse(ROOT / "feeds" / "zbozidostupnost.xml").getroot()
    for item in root.findall("item"):
        code = (item.get("id") or "").strip()
        stock[code] = {
            "qty": int(text(item, "stock_quantity") or "0"),
            "availability": text(item, "availability").lower(),
            "ean": text(item, "ean"),
        }
    return stock


def load_b2b():
    rows = json.loads((ROOT / "feeds" / "avinox-b2b.json").read_text())
    return {row["code"].strip(): row for row in rows}


def build_rows():
    items = list(ET.parse(ROOT / "feeds" / "avinox.xml").getroot().findall("SHOPITEM"))
    stock = load_stock()
    b2b = load_b2b()
    rows = []
    missing_b2b = []
    price_gaps = []
    for item in items:
        productno = text(item, "PRODUCTNO")
        bag = params_of(item)
        base, frame_size = split_size(productno)
        if not frame_size:
            raise SystemExit(f"Mărime lipsă: {productno}")
        model_name = as_model(base)
        family = as_model(text(item, "PRODUCT"))
        color = first(bag, "Colour") or "standard"
        category = first(bag, "E-bike category")
        discipline = DISCIPLINE.get(category, "E-BIKE")
        list_price = int(float(text(item, "PRICE_VAT")))
        live = b2b.get(productno)
        if live is None:
            missing_b2b.append(productno)
        scraped = parse_eu_price(live.get("b2b")) if live else None
        formula = purchase_from_list(list_price)
        purchase = scraped if scraped is not None else formula
        if scraped is not None and abs(scraped - formula) > 0.05:
            price_gaps.append((productno, scraped, formula))

        feed = stock.get(productno, {})
        on_hand = warehouse_qty(live)
        windows = production_windows(live)
        feed_state = feed.get("availability", "")
        if on_hand and on_hand > 0:
            qty = on_hand
            status = "În stoc. Livrare în 5-6 zile."
        elif feed_state == "in stock" and feed.get("qty", 0) > 0 and not windows:
            qty = feed["qty"]
            status = "În stoc. Livrare în 5-6 zile."
        elif feed_state == "sold out" and not windows:
            qty = 0
            status = "Epuizat."
        else:
            qty = 0
            extra = " ".join(windows)
            status = "Precomandă." + (f" {extra}" if extra else "")

        performance = first(bag, "Engine performance")
        location = first(bag, "Engine location")
        motor = performance
        if location:
            motor = join_unique([performance, location])
        front_brake = first(bag, "front brake")
        rear_brake = first(bag, "rear brake")
        brakes = join_unique([
            f"față: {front_brake}" if front_brake else "",
            f"spate: {rear_brake}" if rear_brake else "",
        ])
        crankset = join_unique([
            first(bag, "converter and cranks"),
            first(bag, "cranks"),
            first(bag, "Front rosette"),
        ])
        chain = join_unique([first(bag, "chain"), first(bag, "Belt")])
        cassette = join_unique([first(bag, "cassette / multiwheel"), first(bag, "Rear rosette")])
        shifters = first(bag, "shifting")
        if first(bag, "Automatic shifting") in {"1", "yes", "true"}:
            shifters = join_unique([shifters, "schimbare automată"])
        seatpost = join_unique([first(bag, "seat post"), first(bag, "seat post clamp")])
        wheels = join_unique([
            f"față: {first(bag, 'Tangled front wheel')}" if first(bag, "Tangled front wheel") else "",
            f"spate: {first(bag, 'Tangled rear wheel')}" if first(bag, "Tangled rear wheel") else "",
        ])
        tyre = first(bag, "tires")
        tyre_front = first(bag, "Front tire") or tyre
        tyre_rear = first(bag, "Rear tire") or tyre
        rims = first(bag, "Rims")
        wheel = first(bag, "Wheel diameter")

        description_html = paragraphs_html(unescape(text(item, "DESCRIPTION")))
        ean = text(item, "EAN") or feed.get("ean", "")

        urls = image_urls(item)
        rows.append({
            "category": "E-BIKE",
            "discipline": discipline,
            "brand": "Crussis",
            "family_name": family,
            "model_name": model_name,
            "model_color": color,
            "frame_size": frame_size,
            "sku": productno,
            "list_price_eur": list_price,
            "purchase_price_eur": purchase,
            "stock": qty,
            "stock_status": status,
            "description_html": description_html,
            "description_html_en": description_html,
            "display": first(bag, "Display"),
            "ean": ean,
            "notes": "",
            "image_files": "",
            "frame": first(bag, "Frame"),
            "fork": first(bag, "Fork"),
            "shock": first(bag, "rear shock absorber"),
            "headset": first(bag, "head composition"),
            "stem": first(bag, "stem"),
            "handlebar": first(bag, "handlebars"),
            "grips": first(bag, "grips"),
            "crankset": crankset,
            "rear_derailleur": first(bag, "shifter"),
            "shifters_levers": shifters,
            "brakes": brakes,
            "cassette": cassette,
            "chain": chain,
            "rim_front": rims,
            "rim_rear": rims,
            "hub_front": first(bag, "front hub"),
            "hub_rear": first(bag, "rear hub"),
            "wheels": wheels,
            "tyre_front": tyre_front,
            "tyre_rear": tyre_rear,
            "seatpost": seatpost,
            "saddle": first(bag, "seat"),
            "pedals": first(bag, "pedals"),
            "motor": motor,
            "motor_brand": "Avinox",
            "motor_model": first(bag, "Engine"),
            "battery": battery_text(bag),
            "battery_charger": first(bag, "charger"),
            "motor_power_max_w": watts(performance),
            "wheel_size_front": wheel,
            "wheel_size_rear": wheel,
            "_urls": urls,
            "_feed": feed_state,
        })

    rows.sort(key=lambda row: (
        row["family_name"],
        row["model_name"],
        row["model_color"],
        SIZE_ORDER.get(row["frame_size"], 50),
        row["frame_size"],
        row["sku"],
    ))
    return rows, missing_b2b, price_gaps, set(b2b) - {row["sku"] for row in rows}


def add_portal_sizes(rows, b2b):
    """Mărimi care există în portal, dar nu în fișa publică, când modelul e deja complet."""
    by_model = defaultdict(list)
    for row in rows:
        by_model[row["model_name"]].append(row)
    known = {row["sku"] for row in rows}
    added = []
    for code, live in b2b.items():
        if code in known:
            continue
        base, frame_size = split_size(code)
        if not frame_size:
            continue
        siblings = by_model.get(as_model(base)) or []
        colors = {row["model_color"] for row in siblings}
        if len(colors) != 1:
            continue
        source = siblings[0]
        row = dict(source)
        row["_urls"] = list(source["_urls"])
        row["frame_size"] = frame_size
        row["sku"] = code
        scraped = parse_eu_price(live.get("b2b"))
        if scraped is not None:
            row["purchase_price_eur"] = scraped
        windows = production_windows(live)
        on_hand = warehouse_qty(live)
        if on_hand and on_hand > 0:
            row["stock"] = on_hand
            row["stock_status"] = "În stoc. Livrare în 5-6 zile."
        else:
            row["stock"] = 0
            extra = " ".join(windows)
            row["stock_status"] = "Precomandă." + (f" {extra}" if extra else "")
        row["notes"] = ""
        row["ean"] = ""
        rows.append(row)
        by_model[row["model_name"]].append(row)
        added.append(code)
    rows.sort(key=lambda row: (
        row["family_name"],
        row["model_name"],
        row["model_color"],
        SIZE_ORDER.get(row["frame_size"], 50),
        row["frame_size"],
        row["sku"],
    ))
    return added


def assign_images(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[(row["model_name"], row["model_color"])].append(row["_urls"])
    paths = {}
    for key, lists in grouped.items():
        best = max(lists, key=len)
        model, color = key
        folder = Path("Avinox images") / safe_part(model) / safe_part(color)
        files = []
        for index, url in enumerate(best, start=1):
            name = f"{index:02d}{extension_of(url)}"
            files.append((url, folder / name))
        rel = ";".join(str(path).replace("\\", "/") for _, path in files)
        paths[key] = (rel, files)
    jobs = []
    seen = set()
    for key, (rel, files) in paths.items():
        for row in rows:
            if (row["model_name"], row["model_color"]) == key:
                row["image_files"] = rel
        for url, rel_path in files:
            dest = ROOT / rel_path
            if dest in seen:
                continue
            seen.add(dest)
            jobs.append((url, dest))
    return jobs


def download(url, dest):
    if dest.exists() and dest.stat().st_size > 5000:
        return "skip", url
    dest.parent.mkdir(parents=True, exist_ok=True)
    request = Request(url, headers={"User-Agent": "BikeVerse-catalog/1.0"})
    last_error = None
    for _ in range(3):
        try:
            with urlopen(request, timeout=40) as response:
                data = response.read()
            if len(data) < 1000:
                raise RuntimeError(f"fișier prea mic ({len(data)} bytes)")
            dest.write_bytes(data)
            return "ok", url
        except Exception as error:
            last_error = error
    return "fail", f"{url} :: {last_error}"


def write_files(rows):
    csv_path = ROOT / "avinox-bikeverse.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            out = dict(row)
            out["purchase_price_eur"] = f"{float(row['purchase_price_eur']):.2f}"
            writer.writerow(out)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Avinox"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        sheet.append([row.get(column, "") for column in COLUMNS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"
    price_col = COLUMNS.index("purchase_price_eur") + 1
    for cell in sheet.iter_cols(min_col=price_col, max_col=price_col, min_row=2):
        for item in cell:
            item.number_format = "0.00"
    widths = {
        "model_name": 36, "family_name": 28, "notes": 60, "image_files": 40,
        "stock_status": 42, "sku": 34, "frame": 28, "fork": 36, "motor": 42,
    }
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 18)
    xlsx_path = ROOT / "crussis-avinox-bikeverse.xlsx"
    workbook.save(xlsx_path)
    return csv_path, xlsx_path


def main():
    rows, missing_b2b, price_gaps, extra_b2b = build_rows()
    b2b = load_b2b()
    added = add_portal_sizes(rows, b2b)
    extra_b2b -= set(added)
    print("mărimi adăugate din portal:", ", ".join(added) or "niciuna")
    jobs = assign_images(rows)
    print(f"variante {len(rows)} | modele {len({row['model_name'] for row in rows})} | poze de descărcat {len(jobs)}")
    failed = []
    done = 0
    with ThreadPoolExecutor(max_workers=12) as pool:
        futures = [pool.submit(download, url, dest) for url, dest in jobs]
        for future in as_completed(futures):
            state, info = future.result()
            done += 1
            if state == "fail":
                failed.append(info)
            if done % 100 == 0:
                print(f"poze {done}/{len(jobs)}")
    in_stock = [row for row in rows if row["stock"] > 0]
    preorder = [row for row in rows if row["stock_status"].startswith("Precomandă")]
    sold = [row for row in rows if row["stock_status"].startswith("Epuizat")]
    print(f"în stoc {len(in_stock)} | precomandă {len(preorder)} | epuizat {len(sold)}")
    print("STOC:")
    for row in in_stock:
        print(f"  {row['stock']:>3}  {row['sku']}")
    print(f"fără preț B2B (folosit calculul): {len(missing_b2b)}")
    print(f"preț B2B diferit de formula cu peste 5 cenți: {len(price_gaps)}")
    for gap in price_gaps[:8]:
        print("  ", gap)
    print(f"coduri B2B care nu sunt în XML: {len(extra_b2b)}")
    for code in sorted(extra_b2b)[:30]:
        print("  ", code)
    print(f"poze eșuate: {len(failed)}")
    for info in failed[:10]:
        print("  ", info)
    csv_path, xlsx_path = write_files(rows)
    missing_files = 0
    for row in rows:
        for rel in row["image_files"].split(";"):
            if rel and not (ROOT / rel).exists():
                missing_files += 1
    print("lipsesc fișiere din căi:", missing_files)
    print(csv_path)
    print(xlsx_path)


if __name__ == "__main__":
    main()
